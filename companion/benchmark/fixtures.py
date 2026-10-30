"""Synthetic timeline. Stable Office ZIP timestamps and invariant PDF bytes."""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import re

DAY0=datetime(2026,1,1,tzinfo=timezone.utc).timestamp()


def stable_zip(file):
    src=BytesIO(file.read_bytes());out=BytesIO()
    with ZipFile(src) as z, ZipFile(out,'w') as dest:
        for info in z.infolist():
            info.date_time=(2000,1,1,0,0,0)
            data=z.read(info.filename)
            if info.filename == 'docProps/core.xml':
                data=re.sub(rb'(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)',rb'\g<1>2000-01-01T00:00:00Z\g<2>',data)
            dest.writestr(info,data)
    file.write_bytes(out.getvalue())


def make_corpus(root):
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas
    root=Path(root);star=root/'星河';star.mkdir(parents=True)
    (star/'project.txt').write_text('项目：星河\n负责人：小林\n预算：10万\n状态：立项\n下一步：整理访谈材料\n',encoding='utf-8')
    (star/'plan.md').write_text('# 星河项目\n## 计划\n展示日期：10月8日\n主题：节能本地AI\n待办：完成重启记忆演示\n',encoding='utf-8')
    word=Document();word.add_heading('星河访谈',0);word.add_paragraph('项目：星河');word.add_paragraph('访谈对象：李老师');word.add_paragraph('访谈状态：尚未完成');word.add_paragraph('待办：整理访谈材料和展示讲稿');word.save(star/'report.docx')
    ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='星河课程展示';slide.placeholders[1].text='主题：节能本地AI\n下一步：整理访谈材料';slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='待办';slide.placeholders[1].text='重启记忆演示尚未完成\n展示日期：10月8日';ppt.save(star/'slides.pptx')
    book=Workbook();sheet=book.active;sheet.title='预算';sheet.append(['项目','预算','版本']);sheet.append(['星河','15万','Day 3']);sheet=book.create_sheet('待办');sheet.append(['展示讲稿','尚未完成']);book.save(star/'budget.xlsx')
    pdf=canvas.Canvas(str(star/'brief.pdf'),invariant=1);pdf.drawString(40,780,'Project: Xinghe');pdf.drawString(40,750,'Deadline: October 8');pdf.showPage();pdf.drawString(40,780,'Owner: Lin');pdf.drawString(40,750,'Status: interviews pending, not completed.');pdf.save()
    for file in star.iterdir():
        if file.suffix in {'.docx','.pptx','.xlsx'}:stable_zip(file)
    other=root/'海盐';other.mkdir();(other/'unrelated.txt').write_text('项目：海盐\n预算：99万\n负责人：小陈\n',encoding='utf-8')
    conflict=root/'北辰';conflict.mkdir();(conflict/'budget-old.md').write_text('项目：北辰\n预算：20万\n来源：初稿\n',encoding='utf-8');(conflict/'budget-new.txt').write_text('项目：北辰\n预算：25万\n来源：更新草案，尚未确认\n',encoding='utf-8')
    distract=root/'素材';distract.mkdir()
    for i in range(24):
        (distract/f'noise-{i:02}.txt').write_text(('星河项目 普通保存 文档 项目预算 进度 文件 资料\n'*3)+f'无预算数字；索引素材编号{i}\n',encoding='utf-8')
    # One rare exact key tests retrieval amid common-topic distractors.
    (star/'rare.md').write_text('# 星河资料\n'+'项目资料\n'*240+'特别标记：jadegreen\n供应商：青松印务\n',encoding='utf-8')
    return sorted(root.rglob('*.*'))
