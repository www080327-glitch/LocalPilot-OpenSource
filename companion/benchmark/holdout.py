"""Repeatable frozen V0.2 diagnostics with raw/constrained/product attribution."""
import argparse
import csv
from dataclasses import asdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import statistics
import tempfile
import time

from companion.benchmark.fixtures import stable_zip
from companion.benchmark.runner import canonical_context, digest, model_identity
from companion.benchmark.scoring import confusion, score
from companion.config import Config
from companion.model.structured import decide, parse
from companion.runtime.routing import route_messages
from companion.runtime.service import Companion, SYSTEM
from companion.triggers.active import messages as active_messages

HERE = Path(__file__).parent


def load_holdout():
    sha = (HERE / "holdout.sha256").read_text().strip()
    if digest(HERE / "holdout.json") != sha:
        raise ValueError("Frozen V0.2 holdout hash mismatch")
    return json.loads((HERE / "holdout.json").read_text()), sha


def build_fixture(root, model_path="", backend="disabled"):
    from docx import Document
    from openpyxl import Workbook
    from reportlab.pdfgen.canvas import Canvas
    suite, _ = load_holdout()
    docs = root / "documents"
    primary, other = docs / "松岚", docs / "赤湾"
    primary.mkdir(parents=True)
    other.mkdir()
    (primary / "budget.txt").write_text("项目：松岚\n预算：12万\n", encoding="utf-8")
    (primary / "deadline.md").write_text("项目：松岚\n截止日期：11月16日\n", encoding="utf-8")
    (primary / "long.md").write_text("项目：松岚\n供应商第一次核对：待审\n" + "普通归档资料\n" * 220 +
                                    "供应商第二次核对：银杉印刷\n最终验收：通过\n", encoding="utf-8")
    doc = Document()
    for text in ("项目：松岚", "阶段：准备", "负责人：小陶"):
        doc.add_paragraph(text)
    doc.save(primary / "notes.docx")
    stable_zip(primary / "notes.docx")
    book = Workbook()
    book.active.title = "Costs"
    book.active.append(["项目", "预算", "数量"])
    book.active.append(["松岚", 170000, 170000])
    book.save(primary / "budget.xlsx")
    stable_zip(primary / "budget.xlsx")
    canvas = Canvas(str(primary / "report.pdf"), invariant=1)
    canvas.drawString(70, 700, "Overview")
    canvas.showPage()
    canvas.drawString(70, 700, "Budget: 170000")
    canvas.save()
    for name, amount in (("a.txt", "21万"), ("b.txt", "24万")):
        (other / name).write_text("项目：赤湾\n预算：" + amount + "\n", encoding="utf-8")
    raw = asdict(Config()) | {"data_dir":str(root / "state"), "watch_dirs":[str(docs)],
                             "backend":backend, "model_path":model_path, "project_aliases":suite["aliases"],
                             "dnd":False, "quiet_start_hour":0, "quiet_end_hour":0, "notifications":False}
    filename = root / "config.json"
    filename.write_text(json.dumps(raw, ensure_ascii=False))
    obj = Companion(Config.load(filename))
    obj.scan()
    old = next(r for r in obj.store.records(kinds=["document"]) if r["source_reference"].endswith("budget.txt"))
    (primary / "budget.txt").write_text("项目：松岚\n预算：17万\n", encoding="utf-8")
    obj.process_file(str(primary / "budget.txt"), "refresh")
    obj.store.set("benchmark_old_memory_id", old["id"])
    return obj, docs


def close(obj):
    obj.model.unload()
    obj.channel.server.server_close()
    obj.store.close()
    obj.lock.close()


def source_claim(obj, case):
    docs = obj.store.records(10000, ["document"])
    record = next((r for r in docs if Path(r["source_reference"]).name == case["file"] and obj.store.is_current(r)), None)
    if case.get("old"):
        record = next(r for r in docs if r["id"] == obj.store.get("benchmark_old_memory_id"))
    if record is None:
        return {"memory_id":"invented", "version_id":"invented", "file_hash":"invented",
                "source_file":str(Path(obj.config.watch_dirs[0]) / case["file"]),
                "location":case["location"], "value":case["value"], "field":case["field"]}
    meta = record["metadata"]
    return {"memory_id":"invented" if case.get("fake_id") else record["id"], "version_id":meta["version_id"],
            "file_hash":meta["sha256"], "source_file":record["source_reference"],
            "location":case["location"], "value":case["value"], "field":case["field"]}


def support_messages(obj, case, claim):
    # Full bounded synthetic text and current/old identity made explicit for the raw judge.
    records = obj.store.records(10000, ["document"])
    context = [{"file":Path(r["source_reference"]).name, "current":obj.store.is_current(r),
                "hash":r["metadata"].get("sha256"), "locations":r["metadata"].get("locations"),
                "text":r["metadata"].get("index_text")} for r in records
               if Path(r["source_reference"]).name == case["file"] or case["id"] == "v-conflict" and Path(r["source_reference"]).name in {"a.txt", "b.txt"}]
    public_claim = {k:v for k,v in claim.items() if k not in {"memory_id", "version_id", "source_file"}}
    public_claim.update(file=case["file"], old=bool(case.get("old")), unknown_memory=bool(case.get("fake_id") or case.get("invented")))
    return [{"role":"system", "content":"判断引用是否被来源直接支持：SUPPORTED/PARTIALLY_SUPPORTED(来源冲突)/UNVERIFIED(无法核对)/CONTRADICTED(错误或旧版本)。只返回JSON，support和confidence(0到1)。不能猜测缺失来源。"},
            {"role":"user", "content":json.dumps({"claim":public_claim,"sources":context}, ensure_ascii=False, sort_keys=True)}]


def prompt_hash(messages):
    return hashlib.sha256(json.dumps(messages, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def metrics(rows):
    result = {}
    for layer in ("raw", "constrained", "product"):
        active = [{"expected":r["expected"], "prediction":r[layer]["prediction"]} for r in rows if r["category"] == "active" and layer in r]
        routes = [{"expected":r["expected"], "prediction":r[layer]["prediction"]} for r in rows if r["category"] == "route" and layer in r]
        sources = [r for r in rows if r["category"] == "source" and layer in r]
        formats = [r[layer]["format_valid"] for r in rows if layer in r and "format_valid" in r[layer]]
        durations = [r[layer].get("latency_seconds", r[layer].get("result", {}).get("latency_seconds")) for r in rows if layer in r]
        durations = [v for v in durations if v is not None]
        peaks = [r[layer]["result"]["mlx_peak_bytes"] for r in rows if layer in r and "mlx_peak_bytes" in r[layer].get("result", {})]
        memory = [r for r in rows if r["category"] == "memory" and layer in r]
        result[layer] = {"active":confusion(active, positive=lambda s:s != "quiet"),
                         "route":confusion(routes, positive=lambda s:s == "agent"),
                         "action_accuracy":sum(r["expected"] == r[layer]["prediction"] for r in rows if r["category"] == "active" and layer in r) / len(active) if active else None,
                         "source_support_accuracy":sum(r["expected"] == r[layer]["prediction"] for r in sources) / len(sources) if sources else None,
                         "schema_validity":statistics.mean(formats) if formats else None,
                         "latency_mean_seconds":statistics.mean(durations) if durations else None,
                         "latency_by_category":{cat:statistics.mean([r[layer].get("latency_seconds", r[layer].get("result", {}).get("latency_seconds")) for r in rows if r["category"] == cat and layer in r and r[layer].get("latency_seconds", r[layer].get("result", {}).get("latency_seconds")) is not None]) for cat in ("active","route","source","memory") if any(r["category"] == cat and layer in r and r[layer].get("latency_seconds", r[layer].get("result", {}).get("latency_seconds")) is not None for r in rows)},
                         "mlx_peak_bytes":max(peaks, default=None),
                         "memory_fact_recall":statistics.mean(r[layer]["scores"]["fact_recall"] for r in memory) if memory else None}
    return result


def run(output, model_path="", offline=False, config_file=None):
    suite, sha = load_holdout()
    out = Path(output).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    base = HERE.parent / ".data" / ("holdout-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ"))
    base.mkdir(parents=True)
    obj, docs = build_fixture(base, model_path, "disabled" if offline else "mlx")
    if config_file:
        adapter_config = Config.load(config_file)
        for key in ("backend", "http_model_endpoint", "http_model_name", "http_model_api_key_env", "http_schema_supported"):
            setattr(obj.config, key, getattr(adapter_config, key))
    report = {"status":"running", "kind":"offline_host" if offline else "real_model",
              "suite_sha256":sha,"suite_version":suite["version"], "generation":suite["generation"],
              "started_utc":datetime.now(timezone.utc).isoformat(),"base_dir":str(base),"rows":[],
              "fixture_sha256":{str(p.relative_to(docs)):digest(p) for p in sorted(docs.rglob("*")) if p.is_file()},
              "source_sha256":{str(p.relative_to(HERE.parent)):digest(p) for p in HERE.parent.rglob("*.py") if not any(x in p.parts for x in (".data", ".venv", "build"))},
              "attribution":"Raw and constrained share prompts. Product candidates include host evidence; rule wins are host wins. Each active case resets notification history to isolate classification; rate controls tested separately."}
    if not offline:
        report["identity"] = model_identity(obj.config)
        report["identity_sha256"] = hashlib.sha256(json.dumps(report["identity"], sort_keys=True).encode()).hexdigest()
    report["pipeline_config"] = {k:v for k,v in asdict(obj.config).items() if k not in {"data_dir", "watch_dirs", "model_path", "backend", "http_model_name", "http_model_endpoint", "http_model_api_key_env"}}
    try:
        for case in suite["cases"]:
            row = {"id":case["id"], "category":case["category"], "expected":case.get("expected")}
            cat = case["category"]
            if cat == "active":
                prompt = active_messages(case["event"])
                schema, key = "active", "action"
            elif cat == "route":
                prompt = route_messages(case["query"])
                schema, key = "route", "route"
            elif cat == "source":
                claim = source_claim(obj, case)
                prompt = support_messages(obj, case, claim)
                schema, key = "support", "support"
                began = time.perf_counter()
                checked = obj.verifier.verify(claim)
                row["product"] = {"prediction":checked["status"], "verification":checked, "origin":"host_source_verifier", "latency_seconds":time.perf_counter()-began}
            else:
                memories = obj.retrieve(case["query"])
                row["retrieved_files"] = sorted({Path(r["source_reference"]).name for r in memories if r["type"] == "document"})
                row["retrieval_recall"] = len(set(case["files"]) & set(row["retrieved_files"])) / len(case["files"])
                prompt = [{"role":"system", "content":SYSTEM}, {"role":"user", "content":"来源：" + json.dumps(canonical_context(memories), ensure_ascii=False, sort_keys=True) + "\n问题：" + case["query"]}]
                schema = key = None
            row["prompt"] = prompt
            row["prompt_sha256"] = prompt_hash(prompt)
            if not offline:
                result = obj.model.complete(prompt, max_tokens=96 if schema else 160)
                if schema:
                    decision = parse(result, schema)
                    fallback = {"active":"quiet", "route":"local", "support":"UNVERIFIED"}[schema]
                    row["raw"] = {"result":result,"format_valid":decision is not None, "prediction":decision[key] if decision else fallback}
                    constrained = decide(obj.model, prompt, schema, repairs=obj.config.structured_repairs)
                    decision = constrained["decision"]
                    row["constrained"] = {"result":constrained["attempts"][-1]["result"],"format_valid":decision is not None,
                                          "prediction":decision[key] if decision else fallback, "attempts":constrained["attempts"]}
                else:
                    row["raw"] = {"result":result,"scores":score(case,result),"verification":obj.verifier.answer(result.get("text", ""), memories)}
            if cat == "active" and not offline:
                obj.store.db.execute("DELETE FROM notices")
                obj.store.db.commit()
                obj.store.set("notice_feedback", {})
                obj.store.set("ignore_streak", 0)
                obj.last_candidate.clear()
                began = time.perf_counter()
                answer = obj.consider_event(case["event"])
                row["product"] = {"prediction":answer["action"], "origin":answer["origin"],
                                  "latency_seconds":time.perf_counter()-began,
                                  "notice_emitted":bool(answer.get("notice_id")), "decision":obj.store.get("last_active_decision")}
            if cat == "route" and not offline:
                began = time.perf_counter()
                answer = obj.ask(case["query"])
                # An abstention is not a successful local response.
                row["product"] = {"prediction":"agent" if answer["status"] == "escalation" else "local" if answer["status"] == "ok" else "unavailable",
                                  "answer":answer, "origin":"host_or_semantic_route", "latency_seconds":time.perf_counter()-began}
            if cat == "memory" and not offline:
                began = time.perf_counter()
                answer = obj.ask(case["query"])
                row["product"] = {"result":{"status":answer["status"], "text":answer["text"]}, "scores":score(case,answer), "answer":answer,
                                  "origin":answer.get("model", {}).get("reason", "model_with_verifier")}
                row["product"]["latency_seconds"] = time.perf_counter()-began
            report["rows"].append(row)
            out.write_text(json.dumps(report, ensure_ascii=False, indent=2))
            print(row["id"], row.get("product", {}).get("prediction", row.get("retrieval_recall")), flush=True)
        report["metrics"] = metrics(report["rows"])
        report["retrieval_recall"] = statistics.mean(r["retrieval_recall"] for r in report["rows"] if "retrieval_recall" in r)
        report["model_stats"] = obj.model.stats()
        report["status"] = "completed"
    except BaseException as exc:
        report.update(status="failed", error=f"{type(exc).__name__}: {exc}")
        raise
    finally:
        close(obj)
        out.write_text(json.dumps(report, ensure_ascii=False, indent=2))
        with out.with_suffix(".csv").open("w", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=["id", "category", "expected", "raw", "constrained", "product", "retrieval_recall"])
            writer.writeheader()
            for row in report["rows"]:
                writer.writerow({k:row.get(k) if k not in {"raw", "constrained", "product"} else row.get(k, {}).get("prediction") for k in writer.fieldnames})
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model-path", default="")
    parser.add_argument("--config")
    parser.add_argument("--offline", action="store_true")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    if not args.offline and not args.model_path and not args.config:
        parser.error("Explicit cached model path or local backend config required")
    run(args.output, args.model_path, args.offline, args.config)


if __name__ == "__main__":
    main()
