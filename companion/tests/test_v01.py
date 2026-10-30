"""Reliability and evaluation controls; mock results never prove model capability."""
from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import sqlite3
import tempfile
import time
import unittest
from unittest.mock import Mock,patch

from companion.config import Config,PrivateData
from companion.memory import MemoryStore,LexicalRetriever,BM25Retriever
from companion.observers.documents import extract,extractive_summary
from companion.runtime.service import Companion
from companion.triggers import NoticePolicy
from companion.benchmark.decisions import parse_decision,guard
from companion.benchmark.runner import load_suite,build_store,retrieval_comparison
from companion.benchmark.scoring import confusion,score


class V01Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name).resolve()
        self.docs=self.root/'星河';self.docs.mkdir()
        self.config=Config(low_power_policy="off", localpilot_endpoint="", data_dir=str(self.root/'state'),watch_dirs=(str(self.docs),),backend='disabled',dnd=True,quiet_start_hour=0,quiet_end_hour=0)
        self.obj=Companion(self.config);self.obj.controls._test_bypass=True;self.store=self.obj.store
    def tearDown(self):
        self.obj.model.unload();self.obj.store.close();self.obj.lock.close();self.obj.channel.server.server_close();self.temp.cleanup()
    def index(self,name,text):
        f=self.docs/name;f.write_text(text);self.obj.process_file(str(f),'initial');return f
    def test_current_budget_and_history(self):
        f=self.index('budget.txt','项目：星河\n预算：10万\n')
        old=self.store.records(kinds=['document'])[0]
        f.write_text('项目：星河\n预算：15万\n');self.obj.process_file(str(f),'modified')
        rows=self.obj.retrieve('星河预算')
        self.assertTrue(any('15万' in r['summary'] for r in rows))
        self.assertFalse(any('10万' in r['summary'] for r in rows if r['type']=='document'))
        self.assertEqual(len(self.store.records(kinds=['document'])),2)
        self.assertFalse(self.store.is_current(old))
        self.assertEqual(self.store.db.execute('SELECT count(*) FROM conflicts').fetchone()[0],1)
        self.assertEqual(self.store.active_conflicts(old['id']),[])
    def test_query_checks_hash_before_event_arrives(self):
        f=self.index('budget.txt','项目：星河\n预算：10万\n');f.write_text('项目：星河\n预算：15万\n')
        rows=self.obj.retrieve('星河预算')
        self.assertTrue(any('15万' in r['summary'] for r in rows))
        self.assertFalse(any('10万' in r['summary'] for r in rows if r['type']=='document'))
    def test_failed_update_cannot_return_old_text(self):
        f=self.index('budget.txt','项目：星河\n预算：10万\n');f.write_text('new content')
        with patch('companion.runtime.service.extract',side_effect=ValueError('corrupt')):
            rows=self.obj.retrieve('星河预算')
        self.assertFalse(any('10万' in r['summary'] for r in rows if r['type']=='document'))
        self.assertIn('corrupt',self.obj.last_error)
    def test_deleted_source_cannot_be_factual(self):
        f=self.index('budget.txt','预算：10万');f.unlink()
        self.assertFalse(any(r['type']=='document' for r in self.obj.retrieve('预算')))
    def test_provenance_fields_and_line(self):
        f=self.index('budget.txt','项目：星河\n预算：15万\n')
        row=next(r for r in self.obj.retrieve('预算') if r['type']=='document');p=row['provenance']
        self.assertEqual(p['source_file'],str(f));self.assertEqual(p['file_hash'],hashlib.sha256(f.read_bytes()).hexdigest())
        self.assertTrue(p['version_id']);self.assertGreater(p['extraction_timestamp'],0)
        self.assertEqual(p['document_type'],'txt');self.assertIn('line 2',p['source_location']);self.assertEqual(p['evidence_class'],'OBSERVED')
    def test_conflicting_current_sources_remain_visible(self):
        self.index('a.txt','项目：星河\n预算：10万\n');self.index('b.txt','项目：星河\n预算：15万\n')
        rows=self.obj.retrieve('预算');docs=[r for r in rows if r['type']=='document']
        self.assertEqual(len(docs),2);self.assertTrue(all(r['conflicts'] for r in docs))
    def test_reflection_cannot_be_promoted_through_derivation(self):
        mid=self.store.add('reflection','预算可能20万','companion',epistemic='reflection')
        with self.assertRaises(ValueError):self.store.add('summary','预算20万','model_summary',epistemic='factual',metadata={'derived_from':[mid]})
    def test_model_source_cannot_be_observed(self):
        with self.assertRaises(ValueError):self.store.add('user_fact','用户去过上海','companion_model',epistemic='factual')
    def test_derived_summary_old_version_is_inactive(self):
        f=self.index('budget.txt','预算：10万')
        meta=self.store.records(kinds=['document'])[0]['metadata']
        self.store.add('summary','预算10万','model_document_summary',str(f),epistemic='inferred',metadata={'version_id':meta['version_id'],'sha256':meta['sha256']})
        f.write_text('预算：15万');self.obj.process_file(str(f),'modified')
        self.assertFalse(any('10万' in r['summary'] for r in self.obj.retrieve('预算') if r['type'] in {'document','summary'}))
    def test_restart_current_and_conflict_persistence(self):
        f=self.index('budget.txt','预算：10万');f.write_text('预算：15万');self.obj.process_file(str(f),'modified')
        filename=self.config.data_dir+'/memory.sqlite3'
        second=MemoryStore(filename)
        try:
            rows=BM25Retriever(second).retrieve('预算');self.assertTrue(any('15万' in r['summary'] for r in rows));self.assertFalse(any('10万' in r['summary'] for r in rows if r['type']=='document'))
            self.assertEqual(second.db.execute('SELECT count(*) FROM conflicts').fetchone()[0],1)
        finally:second.close()
    def test_v0_migration_preserves_ids_and_currentness(self):
        file=self.root/'v0.sqlite3';db=sqlite3.connect(file)
        db.execute('CREATE TABLE memories(id TEXT PRIMARY KEY,type TEXT,timestamp REAL,source TEXT,source_reference TEXT,summary TEXT,confidence REAL,epistemic TEXT,metadata TEXT)')
        db.execute('INSERT INTO memories VALUES(?,?,?,?,?,?,?,?,?)',('old-id','document',1,'document_extractor','/docs/a.txt','预算10万',1,'factual',json.dumps({'sha256':'abc'})));db.commit();db.close()
        migrated=MemoryStore(file)
        try:
            row=migrated.records()[0];self.assertEqual(row['id'],'old-id');self.assertTrue(migrated.is_current(row));self.assertTrue(row['metadata']['version_id'])
        finally:migrated.close()
    def test_clear_includes_versions_conflicts_and_fts(self):
        f=self.index('budget.txt','预算：10万');f.write_text('预算：15万');self.obj.process_file(str(f),'modified');self.store.clear()
        for table in ('memories','source_versions','conflicts','claims','memory_fts'):
            self.assertEqual(self.store.db.execute('SELECT count(*) FROM '+table).fetchone()[0],0)
    def test_mute_and_useful_persist_by_type(self):
        self.config.dnd=False;p=NoticePolicy(self.config,self.store)
        nid=p.record('project_help','1','help');p.respond(nid,'mute')
        self.assertFalse(p.allowed('project_help','future',time.time()+86400))
        p.respond(p.record('greeting','2','hi'),'useful')
        self.assertEqual(self.store.get('notice_feedback')['greeting']['useful'],1)
        second=MemoryStore(self.config.data_dir+'/memory.sqlite3')
        try:self.assertFalse(NoticePolicy(self.config,second).allowed('project_help','future',time.time()+86400))
        finally:second.close()
    def test_useful_does_not_bypass_dnd_or_daily_caps(self):
        from datetime import datetime
        t=datetime(2026,10,2,12).timestamp()
        p=NoticePolicy(self.config,self.store);nid=p.record('project_help','1','help',now=t);p.respond(nid,'useful')
        self.assertFalse(p.allowed('project_help','2',t+10000))
        self.config.dnd=False;self.config.max_notices_per_day=1
        self.assertFalse(p.allowed('other','2',t+10000))
    def test_help_prepares_unavailable_escalation(self):
        nid=self.obj.policy.record('project_help','1','help')
        result=self.obj.command({'command':'notice_action','id':nid,'action':'help','text':'整理当前问题'})
        self.assertEqual(result['status'],'escalation');self.assertEqual(result['escalation']['status'],'unavailable')
    def test_pending_escalation_rejects_changed_version(self):
        f=self.index('budget.txt','预算：10万');self.obj.ask('修改预算文件')
        f.write_text('预算：15万');self.obj.process_file(str(f),'modified')
        result=self.obj.command({'command':'submit_escalation','approved':True})
        self.assertEqual(result['status'],'unavailable');self.assertIn('changed',result['reason'])
    def test_fts_deletion_and_quote_safe(self):
        mid=self.store.add('preference','jadegreen','explicit_user_statement')
        self.assertEqual(BM25Retriever(self.store).retrieve('jadegreen')[0]['id'],mid)
        self.store.db.execute('DELETE FROM memories WHERE id=?',(mid,));self.store.db.commit()
        self.assertEqual(BM25Retriever(self.store).retrieve('jadegreen'),[])
        BM25Retriever(self.store).retrieve('" OR * ; --')
    def test_evaluation_corpus_reproducible_and_retrieval_gain(self):
        suite,_=load_suite();a,_,fa=build_store(self.root/'a');b,_,fb=build_store(self.root/'b')
        try:
            self.assertEqual([hashlib.sha256(f.read_bytes()).hexdigest() for f in fa],[hashlib.sha256(f.read_bytes()).hexdigest() for f in fb])
            result=retrieval_comparison(a,suite);self.assertTrue(result['retain_bm25'])
        finally:a.close();b.close()
    def test_comparison_rejects_fixture_or_prompt_drift(self):
        from companion.benchmark.compare import compare
        a={'status':'completed','suite_sha256':'suite','fixture_sha256':{'a':'one'},'generation':{'temperature':0},'retrieval_backend':'bm25','identity_sha256':'model','identity':{'hardware':'host'},'rows':[{'id':'case','category':'greeting','prompt_sha256':'same'}]}
        b=dict(a,fixture_sha256={'a':'two'})
        self.assertEqual(compare(a,b)['status'],'BLOCKED')
        self.assertEqual(compare(a,a)['status'],'accepted_diagnostic_comparison')
    def test_fixture_core_properties_are_time_invariant(self):
        from companion.benchmark.fixtures import stable_zip
        from io import BytesIO
        from zipfile import ZipFile
        f=self.root/'time.xlsx'
        for timestamp in ('2020-01-01T01:02:03Z','2030-06-02T20:30:40Z'):
            with ZipFile(f,'w') as z:z.writestr('docProps/core.xml','<dcterms:modified xmlns:dcterms="urn:test">'+timestamp+'</dcterms:modified>')
            stable_zip(f)
            value=hashlib.sha256(f.read_bytes()).hexdigest()
            if timestamp.startswith('2020'):first=value
            else:self.assertEqual(first,value)
    def test_format_failure_does_not_become_success(self):
        for result in [{'status':'truncated','text':'{"action":"quiet","confidence":1}'},{'status':'ok','text':'{"action":"offer_help","confidence":true}'},{'status':'ok','text':'I think quiet'}]:self.assertIsNone(parse_decision(result))
    def test_false_positive_weight_and_undefined_precision(self):
        c=confusion([{'expected':'quiet','prediction':'remind'},{'expected':'offer_help','prediction':'quiet'}]);self.assertEqual(c['weighted_error_rate'],1)
        self.assertIsNone(confusion([{'expected':'quiet','prediction':'quiet'}])['precision'])
    def test_disturbance_guard_busy_and_confidence(self):
        self.assertEqual(guard({'busy':True},{'action':'offer_help','confidence':1}),'quiet')
        self.assertEqual(guard({'kind':'error'},{'action':'offer_help','confidence':.8}),'quiet')
    def test_cross_spreadsheet_text_conflict_is_recorded_and_exposed(self):
        from openpyxl import Workbook
        self.index('budget.txt','项目：星河\n预算：10万')
        book=Workbook();book.active.title='预算';book.active.append(['项目','预算']);book.active.append(['星河','15万'])
        f=self.docs/'budget.xlsx';book.save(f);self.obj.process_file(str(f),'initial')
        reply=self.obj.ask('星河预算是多少')
        self.assertIn('冲突',reply['text']);self.assertIn('10万',reply['text']);self.assertIn('15万',reply['text'])
        excel=next(r for r in reply['answer_sources'] if r['source_file'].endswith('.xlsx'))
        self.assertIn('sheet 预算 cell B2',excel['source_location'])
    def test_current_field_answer_is_exact_and_does_not_wake_model(self):
        f=self.index('budget.txt','项目：星河\n预算：10万');f.write_text('项目：星河\n预算：15万')
        self.obj.model=Mock()
        reply=self.obj.ask('星河现在预算多少？')
        self.assertIn('15万',reply['text']);self.assertNotIn('10万',reply['text']);self.obj.model.complete.assert_not_called()
        self.assertEqual(reply['answer_sources'][0]['source_location'],['line 2'])
    def test_exact_field_conflict_does_not_pick_silently(self):
        self.index('a.txt','项目：星河\n预算：10万');self.index('b.txt','项目：星河\n预算：15万')
        reply=self.obj.ask('星河预算是多少')
        self.assertIn('冲突',reply['text']);self.assertIn('10万',reply['text']);self.assertIn('15万',reply['text'])
    def test_exact_lookup_respects_requested_file_and_project(self):
        self.index('a.txt','项目：星河\n预算：10万');self.index('b.txt','项目：星河\n预算：15万');self.index('c.txt','项目：海盐\n预算：99万')
        reply=self.obj.ask('只根据a.txt，预算是多少')
        self.assertIn('10万',reply['text']);self.assertNotIn('15万',reply['text']);self.assertNotIn('99万',reply['text'])
        reply=self.obj.ask('星河预算是多少')
        self.assertIn('冲突',reply['text']);self.assertNotIn('99万',reply['text'])
    def test_old_host_answer_cannot_reenter_current_fact_context(self):
        f=self.index('budget.txt','项目：星河\n预算：10万');self.obj.ask('星河预算是多少？')
        f.write_text('项目：星河\n预算：15万');self.obj.process_file(str(f),'modified')
        rows=self.obj.retrieve('星河预算')
        self.assertFalse(any('10万' in r['summary'] for r in rows))
    def test_model_sentence_trace_never_claims_verified(self):
        from companion.memory.answers import sentence_trace
        self.assertEqual(sentence_trace('预算99万。',[])[0]['verification'],'UNVERIFIED_MODEL_CLAIM')
    def test_explicit_denial_is_not_a_forbidden_positive_claim(self):
        self.assertFalse(score({'category':'memory','forbid':['已经完成']},{'status':'ok','text':'无法确认已经完成。'})['detected_hallucination'])
    def test_sheet_and_slide_locations_follow_relationship_order(self):
        from openpyxl import Workbook
        from pptx import Presentation
        book=Workbook();book.remove(book.active)
        for i in range(12):book.create_sheet('sheet'+str(i+1))['A1']='value'+str(i+1)
        f=self.docs/'many.xlsx';book.save(f)
        doc=extract(f.read_bytes(),'.xlsx')
        self.assertIn('[sheet sheet10]\nA1: value10',doc.text)
        ppt=Presentation()
        for text in ('first','second'):
            slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text=text
        ids=ppt.slides._sldIdLst;ids.insert(0,ids[1]);f=self.docs/'reordered.pptx';ppt.save(f)
        doc=extract(f.read_bytes(),'.pptx');self.assertTrue(doc.text.startswith('[slide 1] second'))
    def test_equivalent_whitespace_is_not_a_fact_failure(self):
        from companion.benchmark.scoring import matches
        self.assertTrue(matches('15万', '15 万'))
        self.assertTrue(matches('10月8|October 8', '10 月 8 日'))
    def test_improved_retrieval_finds_rare_fact_at_document_end(self):
        f=self.index('long.md','# 星河\n'+'项目资料\n'*300+'jadegreen供应商：青松印务\n')
        row=next(r for r in self.obj.retrieve('星河jadegreen供应商') if r['type']=='document')
        self.assertIn('青松印务',row['summary'])
    def test_routing_definition_status_gui_complex(self):
        from companion.runtime.routing import needs_agent
        self.assertFalse(needs_agent('解释一下删除是什么意思'))
        self.assertFalse(needs_agent('你刚才已经替我部署系统了吗？'))
        self.assertTrue(needs_agent('帮我操作GUI打开浏览器'))
        self.assertTrue(needs_agent('综合财务表和合同做风险评估'))
    def test_query_invalidated_original_bytes_can_recover(self):
        f=self.index('budget.txt','预算：10万')
        self.store.invalidate_source(str(f));self.obj.process_file(str(f),'refresh')
        self.assertTrue(any('10万' in r['summary'] for r in self.obj.retrieve('预算') if r['type']=='document'))
    def test_submit_checks_change_without_file_event(self):
        f=self.index('budget.txt','预算：10万');self.obj.ask('修改预算文件')
        f.write_text('预算：15万')
        result=self.obj.command({'command':'submit_escalation','approved':True})
        self.assertEqual(result['status'],'unavailable');self.assertIn('changed',result['reason'])
    def test_truncation_is_zero_factual_recall(self):
        case={'category':'document','must':['15万'],'files':['a.txt']}
        self.assertEqual(score(case,{'status':'truncated','text':'15万'})['fact_recall'],0)

if __name__=='__main__':unittest.main()
