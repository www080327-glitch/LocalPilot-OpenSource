"""Synthetic public documents, generated with actual Office/PDF libraries."""
from pathlib import Path

PROJECT = "星河项目"
MARKER = "XINGHE-V0-20261002"


def make_documents(root):
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas
    root = Path(root) / PROJECT
    root.mkdir(parents=True, exist_ok=True)
    (root / "notes.txt").write_text(f"{PROJECT}：用户最近在准备课程展示，主题是节能本地 AI。唯一标记 {MARKER}。", encoding="utf-8")
    (root / "plan.md").write_text(f"# {PROJECT}\n下一步：完成文档长期记忆演示。展示时间：10月8日。", encoding="utf-8")
    word = Document()
    word.add_heading(PROJECT + "汇报", 0)
    word.add_paragraph("负责人：小林。当前工作：整理访谈材料和展示讲稿。")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "里程碑"
    table.cell(0, 1).text = "完成只读记忆测试"
    word.save(root / "report.docx")
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide.shapes.title.text = PROJECT + "课程展示"
    slide.placeholders[1].text = "内容：节能本地 AI\n待办：文档检索、重启记忆演示"
    ppt.save(root / "slides.pptx")
    book = Workbook()
    sheet = book.active
    sheet.title = "预算"
    sheet.append([PROJECT, "金额", "状态"])
    sheet.append(["打印", 120, "已确定"])
    sheet.append(["展示材料", 80, "待确认"])
    sheet["B4"] = "=SUM(B2:B3)"
    book.save(root / "budget.xlsx")
    pdf = canvas.Canvas(str(root / "brief.pdf"))
    pdf.drawString(50, 780, "Xinghe project: local AI course presentation.")
    pdf.drawString(50, 750, "Milestone: persistent document memory demo. Deadline: October 8.")
    pdf.drawString(50, 720, MARKER)
    pdf.save()
    return sorted(root.iterdir())
